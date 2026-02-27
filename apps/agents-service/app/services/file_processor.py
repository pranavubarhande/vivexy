"""
File Processor
Read and write various file types (CSV, XLSX, PDF, DOCX, MD, TXT).
"""

import io
import json
import os
import tempfile
from enum import Enum
from typing import Any, Dict, Optional, Union

import pandas as pd
from docx import Document
from docx.shared import Pt
from pypdf import PdfReader
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.pdfgen import canvas


class FileType(Enum):
    """Supported file types"""

    CSV = "csv"
    XLSX = "xlsx"
    PDF = "pdf"
    DOCX = "docx"
    MD = "md"
    TXT = "txt"
    JSON = "json"


class FileProcessor:
    """
    Handle reading and writing various file types.
    Supports CSV, XLSX, PDF, DOCX, MD, TXT, JSON.
    """

    @staticmethod
    def get_file_type(filepath: str) -> str:
        """Get file type from extension"""
        ext = filepath.split(".")[-1].lower() if "." in filepath else ""
        return ext

    @staticmethod
    def read_file(filepath: str) -> Dict[str, Any]:
        """
        Read file and return content with metadata.

        Args:
            filepath: Path to the file

        Returns:
            Dict with type, data, and text representation
        """
        ext = FileProcessor.get_file_type(filepath)

        try:
            if ext == "csv":
                df = pd.read_csv(filepath)
                return {
                    "type": "csv",
                    "data": df,
                    "text": df.to_string(),
                    "shape": df.shape,
                    "columns": list(df.columns),
                }

            elif ext == "xlsx" or ext == "xls":
                df = pd.read_excel(filepath)
                return {
                    "type": "xlsx",
                    "data": df,
                    "text": df.to_string(),
                    "shape": df.shape,
                    "columns": list(df.columns),
                }

            elif ext == "pdf":
                reader = PdfReader(filepath)
                text = "\n".join([page.extract_text() or "" for page in reader.pages])
                return {
                    "type": "pdf",
                    "data": text,
                    "text": text,
                    "pages": len(reader.pages),
                }

            elif ext == "docx":
                doc = Document(filepath)
                text = "\n".join([para.text for para in doc.paragraphs])
                return {
                    "type": "docx",
                    "data": text,
                    "text": text,
                    "paragraphs": len(doc.paragraphs),
                }

            elif ext == "md":
                with open(filepath, "r", encoding="utf-8") as f:
                    text = f.read()
                return {"type": "md", "data": text, "text": text}

            elif ext == "json":
                with open(filepath, "r", encoding="utf-8") as f:
                    data = json.load(f)
                return {
                    "type": "json",
                    "data": data,
                    "text": json.dumps(data, indent=2),
                }

            else:  # txt or unknown
                with open(filepath, "r", encoding="utf-8") as f:
                    text = f.read()
                return {"type": "txt", "data": text, "text": text}

        except Exception as e:
            return {
                "type": "error",
                "data": None,
                "text": f"Error reading file: {str(e)}",
                "error": str(e),
            }

    @staticmethod
    def write_csv(
        data: Union[pd.DataFrame, Dict, list], filepath: str
    ) -> Dict[str, Any]:
        """
        Write data to CSV file.

        Args:
            data: DataFrame, dict, or list to write
            filepath: Output file path

        Returns:
            Dict with success status and path
        """
        try:
            if isinstance(data, pd.DataFrame):
                df = data
            elif isinstance(data, dict):
                df = pd.DataFrame(data)
            elif isinstance(data, list):
                df = pd.DataFrame(data)
            else:
                df = pd.DataFrame([data])

            df.to_csv(filepath, index=False)

            return {
                "success": True,
                "filepath": filepath,
                "rows": len(df),
                "columns": list(df.columns),
            }

        except Exception as e:
            return {"success": False, "error": str(e)}

    @staticmethod
    def write_xlsx(
        data: Union[pd.DataFrame, Dict, list], filepath: str, sheet_name: str = "Sheet1"
    ) -> Dict[str, Any]:
        """
        Write data to Excel file.

        Args:
            data: DataFrame, dict, or list to write
            filepath: Output file path
            sheet_name: Name for the sheet

        Returns:
            Dict with success status and path
        """
        try:
            if isinstance(data, pd.DataFrame):
                df = data
            elif isinstance(data, dict):
                df = pd.DataFrame(data)
            elif isinstance(data, list):
                df = pd.DataFrame(data)
            else:
                df = pd.DataFrame([data])

            df.to_excel(filepath, index=False, sheet_name=sheet_name)

            return {
                "success": True,
                "filepath": filepath,
                "rows": len(df),
                "columns": list(df.columns),
            }

        except Exception as e:
            return {"success": False, "error": str(e)}

    @staticmethod
    def write_pdf(
        content: str, filepath: str, title: Optional[str] = None
    ) -> Dict[str, Any]:
        """
        Write text content to PDF file.

        Args:
            content: Text content to write
            filepath: Output file path
            title: Optional document title

        Returns:
            Dict with success status and path
        """
        try:
            c = canvas.Canvas(filepath, pagesize=letter)
            width, height = letter

            # Add title if provided
            y = height - 50
            if title:
                c.setFont("Helvetica-Bold", 16)
                c.drawString(50, y, title)
                y -= 30

            # Write content
            c.setFont("Helvetica", 11)
            lines = content.split("\n")

            for line in lines:
                if y < 50:
                    c.showPage()
                    c.setFont("Helvetica", 11)
                    y = height - 50

                # Handle long lines by wrapping
                max_chars = 90
                while len(line) > max_chars:
                    c.drawString(50, y, line[:max_chars])
                    line = line[max_chars:]
                    y -= 14
                    if y < 50:
                        c.showPage()
                        c.setFont("Helvetica", 11)
                        y = height - 50

                c.drawString(50, y, line)
                y -= 14

            c.save()

            return {"success": True, "filepath": filepath, "pages": 1}  # Approximation

        except Exception as e:
            return {"success": False, "error": str(e)}

    @staticmethod
    def write_docx(
        content: str, filepath: str, title: Optional[str] = None
    ) -> Dict[str, Any]:
        """
        Write text content to Word document.

        Args:
            content: Text content to write
            filepath: Output file path
            title: Optional document title

        Returns:
            Dict with success status and path
        """
        try:
            doc = Document()

            # Add title if provided
            if title:
                heading = doc.add_heading(title, level=0)

            # Add content paragraphs
            for para in content.split("\n"):
                if para.strip():
                    p = doc.add_paragraph(para)

            doc.save(filepath)

            return {"success": True, "filepath": filepath}

        except Exception as e:
            return {"success": False, "error": str(e)}

    @staticmethod
    def write_md(content: str, filepath: str) -> Dict[str, Any]:
        """
        Write content to Markdown file.

        Args:
            content: Markdown content
            filepath: Output file path

        Returns:
            Dict with success status and path
        """
        try:
            with open(filepath, "w", encoding="utf-8") as f:
                f.write(content)

            return {"success": True, "filepath": filepath}

        except Exception as e:
            return {"success": False, "error": str(e)}

    @staticmethod
    def write_json(data: Any, filepath: str, indent: int = 2) -> Dict[str, Any]:
        """
        Write data to JSON file.

        Args:
            data: Data to serialize
            filepath: Output file path
            indent: JSON indentation

        Returns:
            Dict with success status and path
        """
        try:
            with open(filepath, "w", encoding="utf-8") as f:
                json.dump(data, f, indent=indent, default=str)

            return {"success": True, "filepath": filepath}

        except Exception as e:
            return {"success": False, "error": str(e)}

    @staticmethod
    def write_txt(content: str, filepath: str) -> Dict[str, Any]:
        """
        Write content to text file.

        Args:
            content: Text content
            filepath: Output file path

        Returns:
            Dict with success status and path
        """
        try:
            with open(filepath, "w", encoding="utf-8") as f:
                f.write(content)

            return {"success": True, "filepath": filepath}

        except Exception as e:
            return {"success": False, "error": str(e)}

    @staticmethod
    def analyze_data(filepath: str) -> Dict[str, Any]:
        """
        Perform statistical analysis on tabular data.

        Args:
            filepath: Path to CSV or XLSX file

        Returns:
            Dict with analysis results
        """
        try:
            file_data = FileProcessor.read_file(filepath)

            if file_data.get("type") not in ["csv", "xlsx"]:
                return {
                    "success": False,
                    "error": "File must be CSV or XLSX for data analysis",
                }

            df = file_data["data"]

            # Basic statistics
            stats = {}

            # Numeric columns
            numeric_cols = df.select_dtypes(include=["number"]).columns.tolist()
            if numeric_cols:
                stats["numeric_summary"] = df[numeric_cols].describe().to_dict()

            # Categorical columns
            cat_cols = df.select_dtypes(include=["object"]).columns.tolist()
            if cat_cols:
                stats["categorical_summary"] = {
                    col: {
                        "unique_values": df[col].nunique(),
                        "top_values": df[col].value_counts().head(5).to_dict(),
                    }
                    for col in cat_cols
                }

            # Missing values
            stats["missing_values"] = df.isnull().sum().to_dict()

            # Shape
            stats["shape"] = {"rows": df.shape[0], "columns": df.shape[1]}

            return {
                "success": True,
                "statistics": stats,
                "columns": list(df.columns),
                "dtypes": df.dtypes.astype(str).to_dict(),
            }

        except Exception as e:
            return {"success": False, "error": str(e)}
