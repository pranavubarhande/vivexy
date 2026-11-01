import email.message
import os
import shutil
import tempfile

import pytest
from docx import Document as DocxDocument
from PIL import Image, ImageDraw, ImageFont
from weasyprint import HTML

from agents_core.services.document_parser import DocumentParseError, parse_document


@pytest.fixture
def workdir():
    path = tempfile.mkdtemp()
    yield path
    shutil.rmtree(path, ignore_errors=True)


def test_parse_text_is_one_unpaginated_page(workdir):
    path = os.path.join(workdir, "notes.txt")
    with open(path, "w") as handle:
        handle.write("Line one\r\nLine two   \n\n")

    parsed = parse_document(path, "txt")

    assert parsed.paginated is False
    assert len(parsed.pages) == 1
    assert parsed.pages[0].text == "Line one\nLine two"


def test_parse_docx_includes_paragraphs_and_tables(workdir):
    path = os.path.join(workdir, "contract.docx")
    docx = DocxDocument()
    docx.add_paragraph("This Agreement is governed by New York law.")
    table = docx.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Term"
    table.rows[0].cells[1].text = "36 months"
    docx.save(path)

    parsed = parse_document(path, "docx")

    assert "governed by New York law" in parsed.pages[0].text
    assert "Term | 36 months" in parsed.pages[0].text


def test_parse_eml_keeps_headers_body_and_attachment_names(workdir):
    message = email.message.EmailMessage()
    message["From"] = "counsel@example.com"
    message["To"] = "client@example.com"
    message["Subject"] = "Settlement offer"
    message.set_content("plain fallback")
    message.add_alternative("<p>We propose <b>$250,000</b>.</p>", subtype="html")
    message.add_attachment(
        b"%PDF", maintype="application", subtype="pdf", filename="offer.pdf"
    )
    path = os.path.join(workdir, "offer.eml")
    with open(path, "wb") as handle:
        handle.write(bytes(message))

    text = parse_document(path, "eml").pages[0].text

    assert "Subject: Settlement offer" in text
    assert "From: counsel@example.com" in text
    assert "plain fallback" in text
    assert "Attachments: offer.pdf" in text


def test_parse_pdf_returns_one_page_per_pdf_page(workdir):
    path = os.path.join(workdir, "brief.pdf")
    HTML(
        string=(
            "<p>First page argues the motion should be granted.</p>"
            "<p style='page-break-before: always'>Second page cites the record.</p>"
        )
    ).write_pdf(path)

    parsed = parse_document(path, "pdf")

    assert parsed.paginated is True
    assert [page.page_number for page in parsed.pages] == [1, 2]
    assert "motion should be granted" in parsed.pages[0].text
    assert "cites the record" in parsed.pages[1].text
    assert parsed.ocr_page_count == 0


@pytest.mark.skipif(shutil.which("tesseract") is None, reason="tesseract not installed")
def test_parse_pdf_ocrs_scanned_pages(workdir):
    image = Image.new("RGB", (1700, 400), "white")
    draw = ImageDraw.Draw(image)
    draw.text(
        (60, 140),
        "INDEMNIFICATION CLAUSE",
        fill="black",
        font=ImageFont.load_default(size=96),
    )
    path = os.path.join(workdir, "scan.pdf")
    image.save(path, "PDF", resolution=150)
    stages = []

    parsed = parse_document(path, "pdf", on_ocr_start=lambda: stages.append("ocr"))

    assert parsed.pages[0].ocr is True
    assert "INDEMNIFICATION" in parsed.pages[0].text.upper()
    assert stages == ["ocr"]


def test_parse_rejects_corrupt_pdf(workdir):
    path = os.path.join(workdir, "broken.pdf")
    with open(path, "wb") as handle:
        handle.write(b"not a pdf")

    with pytest.raises(DocumentParseError):
        parse_document(path, "pdf")


def test_parse_rejects_unsupported_extension(workdir):
    with pytest.raises(DocumentParseError):
        parse_document(os.path.join(workdir, "x.exe"), "exe")
