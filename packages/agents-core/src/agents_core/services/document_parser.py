import email
import email.policy
import logging
from dataclasses import dataclass
from html.parser import HTMLParser
from typing import Callable, List, Optional

import pypdfium2 as pdfium
import pytesseract
from docx import Document as DocxDocument

from agents_core.constants.documents import (
    DOCUMENT_OCR_DPI,
    DOCUMENT_OCR_LANGUAGES,
    DOCUMENT_OCR_MIN_TEXT_CHARS,
)

logger = logging.getLogger(__name__)


@dataclass
class ParsedPage:
    page_number: int
    text: str
    ocr: bool = False


@dataclass
class ParsedDocument:
    pages: List[ParsedPage]
    paginated: bool

    @property
    def ocr_page_count(self) -> int:
        return sum(1 for page in self.pages if page.ocr)


class DocumentParseError(Exception):
    pass


class _HTMLTextExtractor(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts: List[str] = []

    def handle_data(self, data):
        self.parts.append(data)

    def text(self) -> str:
        return "".join(self.parts)


def _normalise(text: str) -> str:
    lines = [line.rstrip() for line in text.replace("\r\n", "\n").split("\n")]
    return "\n".join(lines).strip()


def ocr_image(image) -> str:
    return pytesseract.image_to_string(image, lang=DOCUMENT_OCR_LANGUAGES)


def parse_pdf(
    path: str, on_ocr_start: Optional[Callable[[], None]] = None
) -> ParsedDocument:
    try:
        pdf = pdfium.PdfDocument(path)
    except pdfium.PdfiumError as e:
        raise DocumentParseError(f"Could not open PDF: {e}") from e

    pages: List[ParsedPage] = []
    ocr_started = False
    try:
        for index in range(len(pdf)):
            page = pdf[index]
            text_page = page.get_textpage()
            text = _normalise(text_page.get_text_range())
            text_page.close()
            used_ocr = False
            if len(text) < DOCUMENT_OCR_MIN_TEXT_CHARS:
                if not ocr_started and on_ocr_start:
                    on_ocr_start()
                ocr_started = True
                bitmap = page.render(scale=DOCUMENT_OCR_DPI / 72)
                ocr_text = _normalise(ocr_image(bitmap.to_pil()))
                bitmap.close()
                if len(ocr_text) > len(text):
                    text = ocr_text
                    used_ocr = True
            page.close()
            pages.append(ParsedPage(page_number=index + 1, text=text, ocr=used_ocr))
    finally:
        pdf.close()
    return ParsedDocument(pages=pages, paginated=True)


def parse_docx(path: str) -> ParsedDocument:
    try:
        docx = DocxDocument(path)
    except Exception as e:
        raise DocumentParseError(f"Could not open DOCX: {e}") from e
    blocks = [paragraph.text for paragraph in docx.paragraphs]
    for table in docx.tables:
        for row in table.rows:
            blocks.append(" | ".join(cell.text.strip() for cell in row.cells))
    return ParsedDocument(
        pages=[ParsedPage(page_number=1, text=_normalise("\n".join(blocks)))],
        paginated=False,
    )


def parse_text(path: str) -> ParsedDocument:
    with open(path, "rb") as handle:
        raw = handle.read()
    text = raw.decode("utf-8", errors="replace")
    return ParsedDocument(
        pages=[ParsedPage(page_number=1, text=_normalise(text))], paginated=False
    )


def parse_eml(path: str) -> ParsedDocument:
    with open(path, "rb") as handle:
        message = email.message_from_binary_file(handle, policy=email.policy.default)
    header_lines = [
        f"{name}: {message[name]}"
        for name in ("From", "To", "Cc", "Date", "Subject")
        if message[name]
    ]
    body = message.get_body(preferencelist=("plain", "html"))
    body_text = ""
    if body is not None:
        content = body.get_content()
        if body.get_content_type() == "text/html":
            extractor = _HTMLTextExtractor()
            extractor.feed(content)
            content = extractor.text()
        body_text = content
    attachment_names = [
        part.get_filename()
        for part in message.iter_attachments()
        if part.get_filename()
    ]
    if attachment_names:
        header_lines.append(f"Attachments: {', '.join(attachment_names)}")
    text = "\n".join(header_lines) + "\n\n" + body_text
    return ParsedDocument(
        pages=[ParsedPage(page_number=1, text=_normalise(text))], paginated=False
    )


PARSERS = {
    "pdf": parse_pdf,
    "docx": parse_docx,
    "txt": parse_text,
    "md": parse_text,
    "eml": parse_eml,
}


def parse_document(
    path: str, extension: str, on_ocr_start: Optional[Callable[[], None]] = None
) -> ParsedDocument:
    parser = PARSERS.get(extension)
    if parser is None:
        raise DocumentParseError(f"Unsupported file type: .{extension}")
    if parser is parse_pdf:
        return parse_pdf(path, on_ocr_start=on_ocr_start)
    return parser(path)
